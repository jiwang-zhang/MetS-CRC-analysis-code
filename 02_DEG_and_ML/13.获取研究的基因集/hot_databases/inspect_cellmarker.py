# -*- coding: utf-8 -*-
"""探查 Cell_marker_All.xlsx 结构(表名/列名/前几行)，为后续解析交集做准备。"""
import openpyxl, os
XLSX = r"D:\daixie-colorectal\第二部分DEG\10.获取研究的基因集\hot_databases\Cell_marker_All.xlsx"
wb = openpyxl.load_workbook(XLSX, read_only=True)
print("Sheets:", wb.sheetnames)
for sn in wb.sheetnames:
    ws = wb[sn]
    print(f"\n=== {sn} ({ws.max_row} x {ws.max_column}) ===")
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=3, values_only=True)):
        print(f"  row{i}:", row)
    break  # 只看第一个 sheet 结构
